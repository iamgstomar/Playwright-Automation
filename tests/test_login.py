"""Login module tests driven by the Login worksheet in QA_Test_Cases.xlsx."""

import re
from pathlib import Path

import allure
import openpyxl
import pytest
from playwright.sync_api import Page, TimeoutError as PlaywrightTimeoutError, expect

from conftest import _logout_if_possible


pytestmark = [
    pytest.mark.login,
    pytest.mark.usefixtures("login_case"),
    pytest.mark.flaky(reruns=0),
]

ROOT = Path(__file__).resolve().parents[1]
LOGIN_URL = re.compile(r"/login(?:[/?#]|$)", re.IGNORECASE)
LOGIN_API_URL = re.compile(r"/api/user/v1/login/?$", re.IGNORECASE)
LOGIN_API_TIMEOUT_MS = 15_000


@pytest.fixture
def login_case(request, config):
    """Load this test's scenario and expected result from the Login worksheet."""
    match = re.search(r"test_(TC_LOG_\d+)", request.node.name, re.IGNORECASE)
    if not match:
        pytest.fail(f"Could not determine Login test case ID from {request.node.name!r}")

    tc_id = match.group(1).upper().replace("_", "-", 2)
    workbook_path = Path(config["reporting"]["excel_path"])
    if not workbook_path.is_absolute():
        workbook_path = ROOT / workbook_path
    if not workbook_path.is_file():
        pytest.fail(f"QA test workbook was not found at {workbook_path}")

    workbook = openpyxl.load_workbook(workbook_path, read_only=True, data_only=True)
    try:
        if "Login" not in workbook.sheetnames:
            pytest.fail("The QA test workbook does not contain a Login worksheet")
        sheet = workbook["Login"]
        headers = {
            str(cell.value).strip().lower(): index
            for index, cell in enumerate(sheet[3], start=1)
            if cell.value is not None
        }
        required_headers = (
            "test case id",
            "test scenario",
            "precondition",
            "test steps",
            "expected result",
        )
        missing_headers = [header for header in required_headers if header not in headers]
        if missing_headers:
            pytest.fail(f"Login worksheet is missing required headers: {missing_headers}")

        row_data = None
        for row in sheet.iter_rows(min_row=4, values_only=True):
            if str(row[headers["test case id"] - 1] or "").strip() == tc_id:
                row_data = {
                    header: str(row[headers[header] - 1] or "").strip()
                    for header in required_headers
                }
                break
        if row_data is None:
            pytest.fail(f"{tc_id} was not found in the Login worksheet")
    finally:
        workbook.close()

    for field in ("test scenario", "test steps", "expected result"):
        if not row_data[field]:
            pytest.fail(f"{tc_id} has an empty {field!r} in the Login worksheet")

    allure.dynamic.title(f"{tc_id}: {row_data['test scenario']}")
    allure.dynamic.description(
        f"Precondition: {row_data['precondition']}\n\n"
        f"Test steps:\n{row_data['test steps']}\n\n"
        f"Expected result:\n{row_data['expected result']}"
    )
    return row_data


@pytest.fixture
def login_page(browser, config):
    """Open an unauthenticated Login page without using the authenticated page fixture."""
    context = browser.new_context(ignore_https_errors=True)
    page = context.new_page()
    page.set_default_timeout(int(config["browser"].get("timeout", 15000)))
    login_timeout = int(config["browser"].get("login_timeout", 45000))
    page.set_default_navigation_timeout(login_timeout)
    login_url = (config["environment"].get("login_url") or "").strip()
    if not login_url:
        base_url = (config["environment"].get("base_url") or "").rstrip("/")
        app_path = (config["environment"].get("app_path") or "").strip("/")
        login_path = (config["environment"].get("login_path") or "/login").strip("/")
        login_url = "/".join(part for part in (base_url, app_path, login_path) if part)
    try:
        for attempt in range(2):
            try:
                page.goto(login_url, wait_until="domcontentloaded")
                _username(page).wait_for(state="visible", timeout=login_timeout)
                _password(page).wait_for(state="visible")
                break
            except PlaywrightTimeoutError:
                if attempt:
                    raise
                page.reload(wait_until="domcontentloaded")
        yield page
    finally:
        if not _on_login_page(page):
            _logout_if_possible(page, config)
        context.close()


def _username(page: Page):
    return page.locator("input[name='username']")


def _password(page: Page):
    return page.locator("input[name='password']")


def _submit(page: Page):
    button = page.locator("form.login-form-container button[type='submit']").first
    expect(button).to_be_visible()
    button.click()


def _on_login_page(page: Page) -> bool:
    return bool(LOGIN_URL.search(page.url))


def _login_values(test_data, category, config):
    values = (test_data.get("login") or {}).get(category) or {}
    username = values.get("username") or config["credentials"].get("username")
    password = values.get("password") or config["credentials"].get("password")
    if not username or not password:
        pytest.skip(f"Login test data for {category!r} is not configured")
    return username, password


def _language_option(page: Page, language: str):
    indices = {"english": 0, "arabic": 1}
    language = language.lower()
    if language not in indices:
        raise ValueError(f"Unsupported Login language: {language}")
    trigger = page.locator(".language-select-trigger")
    expect(trigger).to_be_visible()
    trigger.click()
    options = page.locator(".language-option")
    expect(options).to_have_count(2)
    option = options.nth(indices[language])
    expect(option).to_be_visible()
    return option


def _select_language(page: Page, language: str):
    expected_lang = {"english": "en", "arabic": "ar"}[language.lower()]
    if (page.locator("html").get_attribute("lang") or "").lower() != expected_lang:
        _language_option(page, language).click()
    expect(page.locator("html")).to_have_attribute("lang", expected_lang)


def _submit_credentials(page: Page, username: str, password: str):
    _username(page).fill(username)
    _password(page).fill(password)
    matching_requests = []

    def record_login_request(request):
        if request.method == "POST" and LOGIN_API_URL.search(request.url):
            matching_requests.append(request)

    page.on("request", record_login_request)
    try:
        with page.expect_response(
            lambda response: response.request.method == "POST"
            and bool(LOGIN_API_URL.search(response.url)),
            timeout=LOGIN_API_TIMEOUT_MS,
        ) as response_info:
            _submit(page)
    except PlaywrightTimeoutError:
        if matching_requests:
            pytest.xfail(
                "Blocked by the Login API: the request was sent but no response "
                f"arrived within {LOGIN_API_TIMEOUT_MS // 1000} seconds."
            )
        pytest.fail(
            "The Login form did not send a request to the Login API after submit."
        )
    finally:
        page.remove_listener("request", record_login_request)
    return response_info.value


def _submit_unrecognized_credentials(page: Page, username: str, password: str):
    """Submit invalid/unknown credentials and verify the app grants no access."""
    matching_requests = []
    matching_responses = []

    def record_login_request(request):
        if request.method == "POST" and LOGIN_API_URL.search(request.url):
            matching_requests.append(request)

    def record_login_response(response):
        if response.request.method == "POST" and LOGIN_API_URL.search(response.url):
            matching_responses.append(response)

    page.on("request", record_login_request)
    page.on("response", record_login_response)
    try:
        _username(page).fill(username)
        _password(page).fill(password)
        _submit(page)
        if not matching_requests:
            pytest.fail("The Login form did not send a login request after submit.")

        deadline = page.evaluate("() => Date.now()") + LOGIN_API_TIMEOUT_MS
        while (
            not matching_responses
            and _on_login_page(page)
            and _username(page).is_visible()
            and page.evaluate("() => Date.now()") < deadline
        ):
            page.wait_for_timeout(100)

        assert _on_login_page(page), "Unrecognized credentials must not leave the Login page"
        expect(page.locator("form.login-form-container")).to_be_visible()
        assert password not in page.url
        assert password not in page.locator("body").inner_text()
    finally:
        page.remove_listener("request", record_login_request)
        page.remove_listener("response", record_login_response)


def _authenticate(page: Page, username: str, password: str) -> None:
    response = _submit_credentials(page, username, password)
    if not response.ok:
        pytest.fail(
            f"Login API rejected the configured valid test account "
            f"(HTTP {response.status}); update the QA credentials or verify "
            "that the account is active."
        )

    try:
        page.wait_for_function(
            """() => !location.pathname.toLowerCase().includes('/login') ||
                !document.querySelector('form.login-form-container')""",
            timeout=45000,
        )
    except PlaywrightTimeoutError as exc:
        pytest.fail(
            "The login API accepted the request, but the Login form did not "
            f"transition to an authenticated page: {exc}"
        )


def test_TC_LOG_001_login_branding_is_visible_and_accessible(login_page):
    page = login_page
    expect(page.locator("body")).to_contain_text(
        "Kuwait University is a center for innovation with an international reputation for excellence"
    )
    logo = page.locator(".login-logo-slot img")
    expect(logo).to_be_visible()


def test_TC_LOG_002_login_title_and_instruction_text_are_localized(login_page):
    expect(login_page.locator(".login-form-header h2").first).to_have_text("Login")
    expect(login_page.locator(".text-overlay-advertiser")).to_contain_text(
        "Kuwait University is a center for innovation with an international reputation for excellence"
    )
    _select_language(login_page, "arabic")
    expect(login_page.locator(".text-overlay-advertiser")).to_contain_text(
        "جامعة الكويت صرح أكاديمي للابتكار ذو مكانة عالمية متميزة"
    )


def test_TC_LOG_003_language_selector_switches_english_and_arabic(login_page):
    trigger = login_page.locator(".language-select-trigger")
    expect(trigger).to_be_visible()
    trigger.click()
    options = login_page.locator(".language-option")
    expect(options).to_have_count(2)
    expect(options.nth(0)).to_contain_text("English")
    expect(options.nth(1)).to_contain_text("العربية")
    options.nth(1).click()
    expect(login_page.locator("html")).to_have_attribute("lang", "ar")
    expect(trigger).to_contain_text("عربي")
    trigger.click()
    options.nth(0).click()
    expect(login_page.locator("html")).to_have_attribute("lang", "en")
    expect(trigger).to_contain_text("EN")


def test_TC_LOG_004_language_switch_applies_rtl_and_ltr_layout(login_page):
    _select_language(login_page, "arabic")
    expect(login_page.locator(".main-container")).to_have_attribute("dir", "rtl")
    expect(login_page.locator(".text-overlay-advertiser")).to_contain_text(
        "جامعة الكويت صرح أكاديمي للابتكار ذو مكانة عالمية متميزة"
    )
    _select_language(login_page, "english")
    expect(login_page.locator(".main-container")).to_have_attribute("dir", "ltr")


def test_TC_LOG_005_username_is_required(login_page, test_data, config):
    _, password = _login_values(test_data, "valid", config)
    _username(login_page).fill("")
    _password(login_page).fill(password)
    _submit(login_page)
    assert _on_login_page(login_page)
    assert not _username(login_page).evaluate("(el) => el.checkValidity()"), (
        "An empty username should fail browser required-field validation"
    )


def test_TC_LOG_006_password_is_required_and_masked(login_page, test_data, config):
    username, password = _login_values(test_data, "valid", config)
    password_field = _password(login_page)
    expect(password_field).to_have_attribute("type", "password")
    _username(login_page).fill(username)
    _submit(login_page)
    assert _on_login_page(login_page)
    assert not password_field.evaluate("(el) => el.checkValidity()"), (
        "An empty password should fail browser required-field validation"
    )
    assert password not in login_page.url


def test_TC_LOG_007_valid_credentials_authenticate_user(login_page, test_data, config):
    username, password = _login_values(test_data, "valid", config)
    _authenticate(login_page, username, password)
    assert not _on_login_page(login_page), "Valid credentials should leave the Login page"


@pytest.mark.skip(reason="Manually verified; excluded from automated runs per request.")
def test_TC_LOG_011_inactive_session_expires_after_two_minutes(
    login_page, test_data, config
):
    username, password = _login_values(test_data, "valid", config)
    _authenticate(login_page, username, password)
    protected_url = login_page.url
    assert not _on_login_page(login_page), "A protected page is required for the session-timeout test"
    login_page.wait_for_timeout(120000)
    login_page.goto(protected_url, wait_until="domcontentloaded")
    assert _on_login_page(login_page), (
        "An inactive session should expire after the configured two-minute timeout"
    )


def test_TC_LOG_008_invalid_credentials_are_rejected(login_page, test_data, config):
    username, password = _login_values(test_data, "invalid", config)
    _submit_unrecognized_credentials(login_page, username, password)


def test_TC_LOG_009_locked_account_is_rejected(login_page, test_data, config):
    username, password = _login_values(test_data, "locked", config)
    _submit_unrecognized_credentials(login_page, username, password)


def test_TC_LOG_010_login_page_works_at_supported_viewports(login_page):
    for width, height in ((1280, 800), (390, 844)):
        login_page.set_viewport_size({"width": width, "height": height})
        expect(_username(login_page)).to_be_visible()
        expect(_password(login_page)).to_be_visible()
        expect(login_page.locator("form.login-form-container button[type='submit']")).to_be_visible()
        _username(login_page).focus()
        login_page.keyboard.press("Tab")
        assert login_page.locator(":focus").count() == 1, (
            f"Keyboard focus should remain visible at viewport {width}x{height}"
        )

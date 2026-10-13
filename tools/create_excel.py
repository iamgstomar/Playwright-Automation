"""
Creates a fresh QA_Test_Cases.xlsx with the correct structure.
Run once:  python tools/create_excel.py
"""
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
import os

HEADERS = ["Test Case ID", "Test Scenario", "Test Steps", "Expected Result",
           "Pre-conditions", "Priority", "Type", "Status",
           "Actual Result", "Tested By", "Date Tested"]

# Sheet name  ->  list of TC IDs that belong to it
SHEETS = {
    "Login":            [f"TC-LOG-{str(i).zfill(3)}" for i in range(1, 18)],
    "Inbox Broadcast":  [f"TC-BC-{str(i).zfill(3)}"  for i in range(1, 23)],
    "Push Broadcast":   [f"TC-BC-{str(i).zfill(3)}"  for i in range(23, 43)],
    "SMS Broadcast":    [f"TC-BC-{str(i).zfill(3)}"  for i in range(43, 63)],
    "Departments":      [f"TC-DIC-{str(i).zfill(3)}" for i in range(1, 12)],
    "MessageTemplate":  [f"TC-DIC-{str(i).zfill(3)}" for i in range(12, 30)],
    "Credit Management":[f"TC-CM-{str(i).zfill(3)}"  for i in range(1, 11)],
    "CreditReports":    [f"TC-CR-{str(i).zfill(3)}"  for i in range(1, 19)],
    "User Management":  [f"TC-UM-{str(i).zfill(3)}"  for i in range(1, 44)],
    "Reports":          [f"TC-RPT-{str(i).zfill(3)}" for i in range(1, 3)],
    "Home-Header":      [f"TC-HDR-{str(i).zfill(3)}" for i in range(1, 16)],
    "General-CrossCutting": [f"TC-GEN-{str(i).zfill(3)}" for i in range(1, 11)],
}

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(bold=True, color="FFFFFF")
ALT_FILL    = PatternFill("solid", fgColor="D6E4F0")


def _write_sheet(ws, tc_ids):
    # Row 1: module title (merged)
    ws.merge_cells("A1:K1")
    ws["A1"] = ws.title
    ws["A1"].font = Font(bold=True, size=14, color="FFFFFF")
    ws["A1"].fill = PatternFill("solid", fgColor="1F4E79")
    ws["A1"].alignment = Alignment(horizontal="center")

    # Row 2: blank spacer
    # Row 3: column headers
    for col, header in enumerate(HEADERS, start=1):
        cell = ws.cell(row=3, column=col, value=header)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center")

    # Rows 4+: one row per TC ID
    for i, tc_id in enumerate(tc_ids, start=4):
        ws.cell(row=i, column=1, value=tc_id)
        ws.cell(row=i, column=8, value="Not Executed")
        if i % 2 == 0:
            for col in range(1, 12):
                ws.cell(row=i, column=col).fill = ALT_FILL

    # Column widths
    widths = [15, 35, 40, 35, 30, 10, 12, 14, 40, 18, 16]
    for col, width in enumerate(widths, start=1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(col)].width = width


def create():
    wb = openpyxl.Workbook()
    wb.remove(wb.active)  # remove default Sheet

    # Module sheets
    for sheet_name, tc_ids in SHEETS.items():
        ws = wb.create_sheet(title=sheet_name)
        _write_sheet(ws, tc_ids)

    # Summary sheet
    ws = wb.create_sheet(title="Summary")
    ws["A1"] = "Module"
    ws["B1"] = "Total"
    ws["C1"] = "Pass"
    ws["D1"] = "Fail"
    ws["E1"] = "Blocked"
    ws["F1"] = "Not Executed"
    for col in range(1, 7):
        ws.cell(row=1, column=col).font = HEADER_FONT
        ws.cell(row=1, column=col).fill = HEADER_FILL
    for i, sheet_name in enumerate(SHEETS.keys(), start=2):
        ws.cell(row=i, column=1, value=sheet_name)

    out = os.path.join(os.path.dirname(__file__), "..", "QA_Test_Cases.xlsx")
    wb.save(out)
    print(f"Created: {os.path.abspath(out)}")
    print(f"Sheets : {', '.join(SHEETS.keys())}, Summary")


if __name__ == "__main__":
    create()

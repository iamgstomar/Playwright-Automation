import openpyxl
wb = openpyxl.load_workbook('QA_Test_Cases.xlsx')
print('Sheets:', wb.sheetnames)
for name in wb.sheetnames:
    if name == 'Summary':
        continue
    sh = wb[name]
    header_row = 3
    headers = [ (cell.value or '') for cell in sh[header_row] ]
    print('\nSheet:', name)
    print('Header row', header_row, headers)
    for r in range(4, min(sh.max_row, 8) + 1):
        rowvals = [ (sh.cell(row=r, column=c).value or '') for c in range(1, min(12, sh.max_column) + 1) ]
        print('Row', r, rowvals)

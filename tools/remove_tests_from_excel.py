"""
Remove specified TC-DIC test rows from QA_Test_Cases.xlsx.

Usage:
    python tools/remove_tests_from_excel.py

This script will remove rows whose Test Case ID (column A) matches any of the
IDs in the `TO_REMOVE` list across all sheets (except the Summary sheet).
It saves the workbook after modification.
"""
import os
import sys
from openpyxl import load_workbook

ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
EXCEL = os.path.join(ROOT, "QA_Test_Cases.xlsx")
TO_REMOVE = [
    "TC-DIC-016",
    "TC-DIC-017",
    "TC-DIC-018",
    "TC-DIC-019",
    "TC-DIC-029",
]


def main():
    if not os.path.exists(EXCEL):
        print("Excel file not found:", EXCEL)
        sys.exit(1)
    wb = load_workbook(EXCEL)
    changed = False
    for sheet in wb.worksheets:
        if sheet.title == "Summary":
            continue
        # collect row indices to delete
        rows_to_delete = []
        for row in range(4, sheet.max_row + 1):
            cell = sheet.cell(row=row, column=1).value
            if cell and str(cell).strip() in TO_REMOVE:
                rows_to_delete.append(row)
        if rows_to_delete:
            # delete from bottom to top
            for r in sorted(rows_to_delete, reverse=True):
                sheet.delete_rows(r, 1)
            print(f"Removed {len(rows_to_delete)} rows from sheet: {sheet.title}")
            changed = True
    if changed:
        bak = EXCEL + ".bak"
        try:
            os.rename(EXCEL, bak)
            print(f"Backed up original to: {bak}")
        except Exception:
            pass
        wb.save(EXCEL)
        print("Saved updated workbook:", EXCEL)
    else:
        print("No matching test IDs found in workbook. No changes made.")


if __name__ == '__main__':
    main()

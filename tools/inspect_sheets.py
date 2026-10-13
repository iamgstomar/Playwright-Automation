import openpyxl
wb = openpyxl.load_workbook('QA_Test_Cases.xlsx')
for sn in wb.sheetnames:
    if sn == 'Summary':
        continue
    ws = wb[sn]
    ids = [(r, ws.cell(r, 1).value) for r in range(4, ws.max_row + 1) if ws.cell(r, 1).value]
    if ids:
        print(f"{sn}: {ids[0][1]} -> {ids[-1][1]}  ({len(ids)} rows, excel rows {ids[0][0]}-{ids[-1][0]})")
    # also print header row col positions
    headers = {ws.cell(3, c).value: c for c in range(1, ws.max_column + 1) if ws.cell(3, c).value}
    print(f"  Headers: {headers}")

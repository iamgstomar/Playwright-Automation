import openpyxl
import re
from pathlib import Path

WK = 'QA_Test_Cases.xlsx'
OUT = Path('tests') / 'test_user_management.py'


def _detect_columns(sheet, max_header_row=6):
    """Detect common column indexes: tcid, scenario, expected, status.

    Scans the first `max_header_row` rows for header-like text.
    Returns dict with 1-based column indexes.
    """
    headers = {}
    keywords = {
        'tcid': ['tc id', 'tcid', 'test case id', 'test case', 'id'],
        'scenario': ['scenario', 'steps', 'test steps', 'description'],
        'expected': ['expected', 'expected result', 'expected outcome', 'expected result(s)'],
        'status': ['status', 'execution status', 'result', 'test status'],
    }

    for r in range(1, max_header_row + 1):
        rowvals = [str((sheet.cell(row=r, column=c).value or '')).strip().lower() for c in range(1, sheet.max_column + 1)]
        for i, h in enumerate(rowvals, start=1):
            if not h:
                continue
            for key, kws in keywords.items():
                for kw in kws:
                    if kw in h:
                        # prefer the first occurrence
                        if key not in headers:
                            headers[key] = i
    return headers


wb = openpyxl.load_workbook(WK)
sheetname = None
for name in wb.sheetnames:
    if name.strip().lower() == 'user management' or ('user' in name.lower() and 'management' in name.lower()):
        sheetname = name
        break

if not sheetname:
    raise SystemExit(f'User Management sheet not found; available: {wb.sheetnames}')

sh = wb[sheetname]
slug_re = re.compile(r'[^a-z0-9_]')


def slug(s):
    s = str(s).lower()
    s = re.sub(r'[^a-z0-9]+', '_', s)
    s = slug_re.sub('', s)
    if not s:
        return 'unnamed'
    return s


cols = _detect_columns(sh)
# fallback to known defaults if detection fails
tc_col = cols.get('tcid', 1)
scenario_col = cols.get('scenario', 4)
expected_col = cols.get('expected', 7)
status_col = cols.get('status', 9)

OUT.parent.mkdir(parents=True, exist_ok=True)
with OUT.open('w', encoding='utf-8') as f:
    f.write('import pytest\n\n')
    f.write(f"# Auto-generated from sheet: {sheetname}\n\n")
    f.write("# Columns detected: tcid=%d scenario=%d expected=%d status=%d\n\n" % (tc_col, scenario_col, expected_col, status_col))

    written = 0
    for r in range(4, sh.max_row + 1):
        tcid = (sh.cell(row=r, column=tc_col).value or '')
        if tcid is None:
            continue
        tcid = str(tcid).strip()
        if not tcid:
            continue
        scenario = str((sh.cell(row=r, column=scenario_col).value or '')).strip()
        expected = str((sh.cell(row=r, column=expected_col).value or '')).strip()
        status = str((sh.cell(row=r, column=status_col).value or '')).strip()
        # Create a function name that preserves the TC pattern so the report hook can map results.
        # Example TC-UM-001 -> test_TC_UM_001
        safe = re.sub(r'[^0-9A-Za-z]+', '_', tcid).upper()
        func = f'test_{safe}'
        f.write('@pytest.mark.users_roles\n')
        f.write(f'def {func}(page, config):\n')
        f.write(f'    """{tcid}: {scenario}\n\n')
        f.write(f'    Expected:\n    {expected}\n')
        f.write(f'    Sheet: {sheetname} | Status: {status}\n')
        f.write('    """\n')
        f.write("    # TODO: implement this test\n")
        f.write("    pytest.fail('Not implemented: convert this skeleton to real Playwright steps')\n\n")
        written += 1

print('Wrote', OUT, 'tests:', written)


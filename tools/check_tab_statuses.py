import openpyxl
TABS = ['SMS Broadcast','Push Broadcast','Inbox Broadcast','Departments','MessageTemplate']
wb = openpyxl.load_workbook('QA_Test_Cases.xlsx')
for t in TABS:
    if t not in wb.sheetnames:
        print(t, 'MISSING')
        continue
    sh = wb[t]
    total = 0
    not_executed = 0
    for r in range(4, sh.max_row+1):
        tcid = (sh.cell(row=r, column=1).value or '').strip()
        if not tcid:
            continue
        total += 1
        status = (sh.cell(row=r, column=9).value or '').strip()
        if status == '' or status.lower() == 'not executed':
            not_executed += 1
    print(t, 'total:', total, 'not_executed:', not_executed)

import re
import pathlib
import openpyxl

WK = 'QA_Test_Cases.xlsx'
OUT_DIR = pathlib.Path('tests') / 'generated'
OUT_DIR.mkdir(parents=True, exist_ok=True)

wb = openpyxl.load_workbook(WK)

slug_re = re.compile(r'[^a-z0-9_]')

def slug(s):
    s = s.lower().replace(' ', '_')
    s = re.sub(r'[-/\\]', '_', s)
    s = slug_re.sub('', s)
    return s

for sheetname in wb.sheetnames:
    if sheetname == 'Summary':
        continue
    sh = wb[sheetname]
    file_name = f"test_{slug(sheetname)}.py"
    out_path = OUT_DIR / file_name
    with out_path.open('w', encoding='utf-8') as f:
        f.write('import pytest\n\n')
        f.write(f"# Auto-generated from sheet: {sheetname}\n\n")

        for r in range(4, sh.max_row + 1):
            tcid = sh.cell(row=r, column=1).value
            if not tcid:
                continue
            scenario = sh.cell(row=r, column=4).value or ''
            expected = sh.cell(row=r, column=7).value or ''
            status = (sh.cell(row=r, column=9).value or '').strip()
            func_name = f'test_{slug(tcid)}'
            f.write(f"@pytest.mark.skip(reason=\"Auto-generated skeleton\")\n")
            f.write(f"def {func_name}():\n")
            f.write(f"    \"\"\"{tcid}: {scenario}\n\n")
            f.write(f"    Expected:\n    {expected}\n")
            f.write(f"    Status in sheet: {status}\n")
            f.write(f"    \"\"\"\n")
            f.write(f"    pytest.skip('Implement this test based on the workbook')\n\n")
    print('Wrote', out_path)

print('Done')

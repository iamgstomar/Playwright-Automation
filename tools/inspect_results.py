import openpyxl
from pathlib import Path

candidates = [Path('QA_Test_Cases.xlsx'), Path('..') / 'QA_Test_Cases.xlsx']
for p in candidates:
    if p.exists():
        print('Found workbook at', p)
        wb = openpyxl.load_workbook(p)
        if 'User Management' in wb.sheetnames:
            sh = wb['User Management']
            # detect status column similar to excel_writer
            header_row = 3
            headers = [str(cell.value).strip().lower() if cell.value is not None else '' for cell in sh[header_row]]
            status_idx = None
            for i,h in enumerate(headers, start=1):
                if 'status' in h or 'result' in h:
                    status_idx = i
                    break
            if status_idx is None:
                status_idx = 8
            print('Using status column', status_idx)
            for r in range(4, min(sh.max_row+1, 4+50)):
                tc = sh.cell(row=r, column=1).value
                status = sh.cell(row=r, column=status_idx).value
                if tc:
                    print(r, tc, '=>', status)
        else:
            print('No User Management sheet in', p)
    else:
        print('Workbook not found at', p)
